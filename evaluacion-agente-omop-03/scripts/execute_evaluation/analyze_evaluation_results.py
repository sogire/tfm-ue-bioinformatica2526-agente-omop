#!/usr/bin/env python3
"""
Script de consolidación y análisis de resultados de evaluación OMOP.
Lee los archivos YAML de evaluación (test) y resultados (result),
los cruza en un pandas DataFrame y exporta el informe consolidado a Excel con timestamp
bajo el directorio 'results/'.
"""

import os
import re
import glob
import yaml
import datetime
import argparse
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def natural_sort_key(s):
    """Clave de ordenación natural para cadenas que contienen números (ej. C01, C02, C10)."""
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', str(s))]

def build_test_dataframe(eval_dir):
    """
    Lee recursivamente los archivos YAML bajo el directorio de evaluación
    y genera el DataFrame 'test' con los campos:
    - source
    - id_evaluacion
    - nivel_complejidad
    - domain_main
    - domain_others
    - pregunta_lenguaje_natural
    - sql_base_gold_standard
    """
    pattern = os.path.join(eval_dir, "**/*.yaml")
    files = glob.glob(pattern, recursive=True)
    rows = []
    
    for fpath in files:
        try:
            with open(fpath, "r", encoding="utf-8") as f:
                data = yaml.safe_load(f)
            if not isinstance(data, dict):
                continue
            eval_id = data.get("id_evaluacion")
            source_val = data.get("source")
            if not source_val:
                if eval_id and "_" in str(eval_id):
                    source_val = str(eval_id).split("_", 1)[0]
                else:
                    source_val = os.path.basename(os.path.dirname(fpath))
                    
            complejidad = data.get("nivel_complejidad")
            domain_main = data.get("domain_main", "")
            domain_others = data.get("domain_others", "")
            pregunta = data.get("pregunta_lenguaje_natural")
            sql_gold = data.get("sql_base_gold_standard", "")
            if eval_id:
                clean_pregunta = re.sub(r"\s+", " ", str(pregunta)).strip() if pregunta else ""
                rows.append({
                    "source": str(source_val).strip(),
                    "id_evaluacion": str(eval_id).strip(),
                    "nivel_complejidad": str(complejidad).strip() if complejidad else "no especificado",
                    "domain_main": str(domain_main).strip() if domain_main else "",
                    "domain_others": str(domain_others).strip() if domain_others else "",
                    "pregunta_lenguaje_natural": clean_pregunta,
                    "sql_base_gold_standard": str(sql_gold).strip() if sql_gold else ""
                })
        except Exception as err:
            print(f"[Aviso] Error leyendo archivo de evaluación {fpath}: {err}")
            
    df_test = pd.DataFrame(rows, columns=[
        "source",
        "id_evaluacion",
        "nivel_complejidad",
        "domain_main",
        "domain_others",
        "pregunta_lenguaje_natural",
        "sql_base_gold_standard"
    ])
    df_test.drop_duplicates(subset=["id_evaluacion"], inplace=True)
    return df_test

def build_result_dataframe(results_dir):
    """
    Lee recursivamente los archivos YAML bajo el directorio de resultados
    y genera el DataFrame 'result' con los campos globales y por turno:
    - id_evaluacion
    - check_conexion
    - turnos_ejecutados
    - validez_sintactica_y_estructural
    - precision_terminologica_vocabulario
    - turno_1_validez_sintactica_y_estructural
    - turno_1_precision_terminologica_vocabulario
    - turno_1_question_improvement
    - turno_2_validez_sintactica_y_estructural
    - turno_2_precision_terminologica_vocabulario
    - turno_2_question_improvement
    - turno_3_validez_sintactica_y_estructural
    - turno_3_precision_terminologica_vocabulario
    - turno_3_question_improvement
    - evaluacion_cualitativa
    - sql_generada_agente
    """
    pattern = os.path.join(results_dir, "**/*.yaml")
    files = glob.glob(pattern, recursive=True)
    rows = []
    
    for fpath in files:
        try:
            with open(fpath, "r", encoding="utf-8") as f:
                data = yaml.safe_load(f)
            if not isinstance(data, dict):
                continue
            eval_id = data.get("id_evaluacion")
            if not eval_id:
                continue
                
            check_conexion = data.get("check_conexion")
            turnos_ejecutados = data.get("turnos_ejecutados", 1)
            sql_agente = data.get("sql_generada_agente", "")
            metrics = data.get("metricas_evaluacion", {}) or {}
            validez = metrics.get("validez_sintactica_y_estructural", 0.0)
            vocab = metrics.get("precision_terminologica_vocabulario", 0.0)
            cualitativa = data.get("evaluacion_cualitativa", "")
            
            # Metadatos de sesión
            session_id = data.get("session_id", "")
            session_file = data.get("session_file", "")
            case_name = os.path.basename(fpath).replace(".yaml", "")
            source_name = os.path.basename(os.path.dirname(fpath))
            if not session_file or session_file.startswith("sessions/"):
                session_file = f"results/{source_name}/sessions/{case_name}.json"
            
            # Desglose de métricas y question_improvement por turno
            metricas_por_turno = metrics.get("metricas_por_turno", {}) or {}
            
            # Turno 1
            t1 = metricas_por_turno.get("turno_1", {}) or {}
            t1_validez = t1.get("validez_sintactica_y_estructural")
            if t1_validez is None and turnos_ejecutados == 1:
                t1_validez = validez
            t1_vocab = t1.get("precision_terminologica_vocabulario")
            if t1_vocab is None and turnos_ejecutados == 1:
                t1_vocab = vocab
            t1_qi = t1.get("question_improvement", "")
            
            # Turno 2
            t2 = metricas_por_turno.get("turno_2", {}) or {}
            t2_validez = t2.get("validez_sintactica_y_estructural")
            t2_vocab = t2.get("precision_terminologica_vocabulario")
            t2_qi = t2.get("question_improvement", "")
            
            # Turno 3
            t3 = metricas_por_turno.get("turno_3", {}) or {}
            t3_validez = t3.get("validez_sintactica_y_estructural")
            t3_vocab = t3.get("precision_terminologica_vocabulario")
            t3_qi = t3.get("question_improvement", "")
            
            # Desglose de Criterios Métrica 1
            detalle_sintaxis = metrics.get("detalle_validez_sintactica", {}) or {}
            c1_1 = detalle_sintaxis.get("criterio_1_1_ejecucion")
            c1_2 = detalle_sintaxis.get("criterio_1_2_similitud_sql_gemini")
            justificacion_gemini = detalle_sintaxis.get("justificacion_gemini", "")
            
            rows.append({
                "id_evaluacion": str(eval_id).strip(),
                "session_id": str(session_id).strip() if session_id else "",
                "session_file": str(session_file).strip() if session_file else "",
                "check_conexion": check_conexion,
                "turnos_ejecutados": int(turnos_ejecutados) if turnos_ejecutados is not None else 1,
                "validez_sintactica_y_estructural": float(validez) if validez is not None else 0.0,
                "criterio_1_1_ejecucion": float(c1_1) if c1_1 is not None else None,
                "criterio_1_2_similitud_sql_gemini": float(c1_2) if c1_2 is not None else None,
                "precision_terminologica_vocabulario": float(vocab) if vocab is not None else 0.0,
                "turno_1_validez_sintactica_y_estructural": float(t1_validez) if t1_validez is not None else None,
                "turno_1_precision_terminologica_vocabulario": float(t1_vocab) if t1_vocab is not None else None,
                "turno_1_question_improvement": str(t1_qi).strip() if t1_qi else "",
                "turno_2_validez_sintactica_y_estructural": float(t2_validez) if t2_validez is not None else None,
                "turno_2_precision_terminologica_vocabulario": float(t2_vocab) if t2_vocab is not None else None,
                "turno_2_question_improvement": str(t2_qi).strip() if t2_qi else "",
                "turno_3_validez_sintactica_y_estructural": float(t3_validez) if t3_validez is not None else None,
                "turno_3_precision_terminologica_vocabulario": float(t3_vocab) if t3_vocab is not None else None,
                "turno_3_question_improvement": str(t3_qi).strip() if t3_qi else "",
                "evaluacion_cualitativa": str(cualitativa).strip() if cualitativa else "",
                "evaluacion_llm_juez": str(justificacion_gemini).strip() if justificacion_gemini else "",
                "sql_generada_agente": str(sql_agente).strip() if sql_agente else ""
            })
        except Exception as err:
            print(f"[Aviso] Error leyendo archivo de resultado {fpath}: {err}")
            
    df_result = pd.DataFrame(rows, columns=[
        "id_evaluacion",
        "session_id",
        "session_file",
        "check_conexion",
        "turnos_ejecutados",
        "validez_sintactica_y_estructural",
        "criterio_1_1_ejecucion",
        "criterio_1_2_similitud_sql_gemini",
        "precision_terminologica_vocabulario",
        "turno_1_validez_sintactica_y_estructural",
        "turno_1_precision_terminologica_vocabulario",
        "turno_1_question_improvement",
        "turno_2_validez_sintactica_y_estructural",
        "turno_2_precision_terminologica_vocabulario",
        "turno_2_question_improvement",
        "turno_3_validez_sintactica_y_estructural",
        "turno_3_precision_terminologica_vocabulario",
        "turno_3_question_improvement",
        "evaluacion_cualitativa",
        "evaluacion_llm_juez",
        "sql_generada_agente"
    ])
    df_result.drop_duplicates(subset=["id_evaluacion"], inplace=True)
    return df_result

def export_to_excel(df, output_path):
    """Exporta el DataFrame a un libro Excel con formateo profesional y autoajuste de columnas."""
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Resultados_Evaluacion")
        ws = writer.sheets["Resultados_Evaluacion"]
        
        # Estilos visuales: Fuente Negra, BackColor: Blanco (sin complicaciones de color)
        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="000000")
        regular_font = Font(name="Arial", size=11, color="000000")
        sql_font = Font(name="Consolas", size=9, color="000000")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        
        for col_idx in range(1, len(df.columns) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = white_fill
            cell.font = header_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
        for row in range(2, len(df) + 2):
            ws.row_dimensions[row].height = 50  # Altura suficiente para multilínea
            for col in range(1, len(df.columns) + 1):
                c = ws.cell(row=row, column=col)
                c.fill = white_fill
                c.font = regular_font
                c.border = thin_border
                col_name = df.columns[col - 1]
                if col_name in [
                    "validez_sintactica_y_estructural",
                    "criterio_1_1_ejecucion",
                    "criterio_1_2_similitud_sql_gemini",
                    "precision_terminologica_vocabulario",
                    "turno_1_validez_sintactica_y_estructural",
                    "turno_1_precision_terminologica_vocabulario",
                    "turno_2_validez_sintactica_y_estructural",
                    "turno_2_precision_terminologica_vocabulario",
                    "turno_3_validez_sintactica_y_estructural",
                    "turno_3_precision_terminologica_vocabulario"
                ]:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    if c.value is not None and not (isinstance(c.value, float) and pd.isna(c.value)):
                        c.number_format = "0.00"
                elif col_name in ["source", "id_evaluacion", "session_id", "session_file", "nivel_complejidad", "domain_main", "check_conexion", "turnos_ejecutados"]:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif col_name in ["domain_others", "session_file"]:
                    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                elif col_name in ["sql_base_gold_standard", "sql_generada_agente"]:
                    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                    c.font = sql_font
                elif col_name in ["turno_1_question_improvement", "turno_2_question_improvement", "turno_3_question_improvement"]:
                    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                else:
                    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                    
        # Ancho óptimo de columnas
        column_widths = {
            "source": 16,
            "id_evaluacion": 22,
            "session_id": 38,
            "session_file": 40,
            "nivel_complejidad": 18,
            "domain_main": 22,
            "domain_others": 28,
            "pregunta_lenguaje_natural": 50,
            "check_conexion": 16,
            "turnos_ejecutados": 18,
            "validez_sintactica_y_estructural": 22,
            "criterio_1_1_ejecucion": 18,
            "criterio_1_2_similitud_sql_gemini": 22,
            "precision_terminologica_vocabulario": 24,
            "turno_1_validez_sintactica_y_estructural": 24,
            "turno_1_precision_terminologica_vocabulario": 26,
            "turno_1_question_improvement": 30,
            "turno_2_validez_sintactica_y_estructural": 24,
            "turno_2_precision_terminologica_vocabulario": 26,
            "turno_2_question_improvement": 45,
            "turno_3_validez_sintactica_y_estructural": 24,
            "turno_3_precision_terminologica_vocabulario": 26,
            "turno_3_question_improvement": 45,
            "evaluacion_cualitativa": 65,
            "evaluacion_llm_juez": 65,
            "sql_base_gold_standard": 60,
            "sql_generada_agente": 60
        }
        for col_idx, col_name in enumerate(df.columns, 1):
            letter = get_column_letter(col_idx)
            ws.column_dimensions[letter].width = column_widths.get(col_name, 20)

def main():
    parser = argparse.ArgumentParser(
        description="Consolidación y exportación a Excel del análisis de evaluación OMOP."
    )
    parser.add_argument(
        "--source", "-s",
        default="querylibrary",
        help="Identificador de la fuente a analizar (por defecto: 'querylibrary', o 'all' para todas las fuentes)"
    )
    parser.add_argument(
        "--eval-dir", "-e",
        default=None,
        help="Directorio base de archivos de prueba Gold Standard (por defecto: evaluation/gold_standard/<source>)"
    )
    parser.add_argument(
        "--results-dir", "-r",
        default=None,
        help="Directorio base de archivos de resultado YAML (por defecto: evaluation/results/<source>)"
    )
    parser.add_argument(
        "--output", "-o",
        help="Ruta o nombre personalizado del archivo Excel de salida (opcional)"
    )
    parser.add_argument(
        "--base-dir", "-b",
        default=None,
        help="Directorio raíz del proyecto (por defecto: auto-detectado)"
    )
    
    args = parser.parse_args()
    
    if args.base_dir:
        base_dir = os.path.abspath(args.base_dir)
    else:
        script_dir = os.path.dirname(os.path.abspath(__file__))
        if os.path.exists(os.path.join(script_dir, "evaluation")):
            base_dir = script_dir
        else:
            base_dir = os.path.abspath(os.path.join(script_dir, "../.."))
            
    source = args.source
    
    if args.eval_dir:
        eval_dir = os.path.abspath(args.eval_dir) if os.path.isabs(args.eval_dir) else os.path.join(base_dir, args.eval_dir)
    elif source and source.lower() != "all":
        eval_dir = os.path.join(base_dir, "evaluation/gold_standard", source)
    else:
        eval_dir = os.path.join(base_dir, "evaluation/gold_standard")
        
    if args.results_dir:
        results_dir = os.path.abspath(args.results_dir) if os.path.isabs(args.results_dir) else os.path.join(base_dir, args.results_dir)
    elif source and source.lower() != "all":
        results_dir = os.path.join(base_dir, "evaluation/results", source)
    else:
        results_dir = os.path.join(base_dir, "evaluation/results")
    
    print(f"Fuente seleccionada: {source}")
    print(f"Cargando archivos de evaluación desde: {eval_dir}")
    df_test = build_test_dataframe(eval_dir)
    print(f"-> Casos definidos en 'test': {len(df_test)}")
    
    print(f"\nCargando archivos de resultados desde: {results_dir}")
    df_result = build_result_dataframe(results_dir)
    print(f"-> Casos evaluados en 'result': {len(df_result)}")
    
    if df_result.empty:
        print("[Aviso] No se encontraron resultados YAML en el directorio especificado.")
        return
        
    # Cruce de dataframes: combinamos las pruebas evaluadas con sus metadatos (complejidad y pregunta)
    df_merged = pd.merge(df_test, df_result, on="id_evaluacion", how="inner")
    
    # Ordenación natural de los casos evaluados
    df_merged["_sort_key"] = df_merged["id_evaluacion"].apply(natural_sort_key)
    df_merged.sort_values(by="_sort_key", inplace=True)
    df_merged.drop(columns=["_sort_key"], inplace=True)
    
    # Columnas finales consolidadas con 'source' en primera posición
    column_order = [
        "source",
        "id_evaluacion",
        "session_id",
        "session_file",
        "nivel_complejidad",
        "domain_main",
        "domain_others",
        "pregunta_lenguaje_natural",
        "check_conexion",
        "turnos_ejecutados",
        "validez_sintactica_y_estructural",
        "criterio_1_1_ejecucion",
        "criterio_1_2_similitud_sql_gemini",
        "precision_terminologica_vocabulario",
        "turno_1_validez_sintactica_y_estructural",
        "turno_1_precision_terminologica_vocabulario",
        "turno_1_question_improvement",
        "turno_2_validez_sintactica_y_estructural",
        "turno_2_precision_terminologica_vocabulario",
        "turno_2_question_improvement",
        "turno_3_validez_sintactica_y_estructural",
        "turno_3_precision_terminologica_vocabulario",
        "turno_3_question_improvement",
        "evaluacion_cualitativa",
        "evaluacion_llm_juez",
        "sql_base_gold_standard",
        "sql_generada_agente"
    ]
    df_merged = df_merged[column_order]
    
    # Generar nombre del archivo con timestamp actual yyyymmddhhmm bajo 'evaluation/results' con prefijo <source>_
    timestamp = datetime.datetime.now().strftime("%Y%m%d%H%M")
    if args.output:
        output_file = args.output if os.path.isabs(args.output) else os.path.join(base_dir, args.output)
    else:
        output_file = os.path.join(base_dir, "evaluation/results", f"{source}_analisis_{timestamp}.xlsx")
        
    os.makedirs(os.path.dirname(output_file), exist_ok=True)
    export_to_excel(df_merged, output_file)
    
    print("\n" + "="*80)
    print(f"INFORME CONSOLIDADO GENERADO: {output_file}")
    print("="*80)
    print(f"Total de casos consolidados: {len(df_merged)}")
    
    # Resumen en consola agrupado por nivel de complejidad
    print("\nResumen promedio por nivel de complejidad:")
    summary = df_merged.groupby("nivel_complejidad")[["validez_sintactica_y_estructural", "precision_terminologica_vocabulario"]].agg(["count", "mean"])
    print(summary)
    
    print("\nDetalle de filas consolidadas:")
    for _, row in df_merged.iterrows():
        dom_info = f"{row['domain_main']}" + (f" (+{row['domain_others']})" if row['domain_others'] else "")
        t_info = f"Turnos: {row.get('turnos_ejecutados', 1)}"
        print(f" • [{row['source']}] {row['id_evaluacion']:<20} | Complejidad: {row['nivel_complejidad']:<11} | Dominio: {dom_info:<26} | {t_info:<10} | Sintaxis: {row['validez_sintactica_y_estructural']:.2f} | Vocabulario: {row['precision_terminologica_vocabulario']:.2f} | Conexión: {row['check_conexion']}")

if __name__ == "__main__":
    main()
